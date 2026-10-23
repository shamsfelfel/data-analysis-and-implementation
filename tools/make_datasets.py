#!/usr/bin/env python3
"""Generate the course's shared sample datasets, deterministically.

    python tools/make_datasets.py

Writes into downloads/ (shipped with the site, linked from the lessons) and
docs/datasets.md (the facts every lesson quotes, computed here from the files
so a lesson can never disagree with the data).

Theory 1  -- rental listings (CSV, semicolon CSV with Arabic, Excel with serial
             dates) and a paginated school-registry API (JSON fixtures)
Practical 1 -- a campus cafeteria's orders for one month (CSV, Excel with a
             Table, and a menu sheet), deliberately a little messy.

The mess is deliberate and documented: duplicates, missing values, spelling
variants, a price typo, refunds. Lessons teach students to find it.
"""
import csv
import datetime as dt
import io
import json
import os
import random

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
T1 = os.path.join(ROOT, "downloads", "t1")
P1 = os.path.join(ROOT, "downloads", "p1")
for d in (T1, P1):
    os.makedirs(d, exist_ok=True)

facts = []


def fact(s):
    facts.append(s)


# ------------------------------------------------------------------ Theory 1: rentals

rnd = random.Random(2305)
CITIES = [("Gaza", "غزة", 420, 85), ("Khan Younis", "خان يونس", 300, 80), ("Rafah", "رفح", 250, 75),
          ("Deir al-Balah", "دير البلح", 270, 78), ("Jabalia", "جباليا", 230, 82)]
GOV = {"Gaza": "Gaza", "Khan Younis": "Khan Younis", "Rafah": "Rafah",
       "Deir al-Balah": "Deir al-Balah", "Jabalia": "North Gaza"}
ADVERT = [
    "Bright flat near the market, second floor, quiet street",
    "Sunny apartment, balcony, newly painted, water tank included",
    "Ground floor, close to the university, suitable for students",
    "Family flat, three bedrooms, big kitchen, generator available",
    "Small studio, ideal for one person, internet ready",
]

rentals = []
base = dt.date(2026, 1, 5)
for i in range(1, 61):
    city_en, city_ar, rent0, area0 = rnd.choice(CITIES)
    rooms = rnd.choice([1, 2, 2, 3, 3, 3, 4, 5])
    area = round(area0 * (0.55 + 0.22 * rooms) * rnd.uniform(0.85, 1.15), 1)
    rent = int(round((rent0 * (0.6 + 0.15 * rooms)) * rnd.uniform(0.85, 1.2) / 5.0) * 5)
    furnished = rnd.choice(["yes", "no", "no"])
    rentals.append({
        "listing_id": 1000 + i,
        "city": city_en, "city_ar": city_ar, "governorate": GOV[city_en],
        "monthly_rent": rent, "area_m2": area, "rooms": rooms, "furnished": furnished,
        "listed_on": (base + dt.timedelta(days=rnd.randint(0, 120))).isoformat(),
        "advert": rnd.choice(ADVERT),
    })
# deliberate gaps
for i in (4, 17, 33, 48):
    rentals[i]["area_m2"] = ""
for i in (9, 25):
    rentals[i]["monthly_rent"] = ""
rentals[41]["rooms"] = ""

cols = ["listing_id", "city", "governorate", "monthly_rent", "area_m2", "rooms", "furnished", "listed_on", "advert"]
with io.open(os.path.join(T1, "rental_listings.csv"), "w", encoding="utf-8", newline="") as f:
    w = csv.DictWriter(f, fieldnames=cols, extrasaction="ignore")
    w.writeheader()
    w.writerows(rentals)

# semicolon, decimal comma, Arabic city names, UTF-8 with BOM (what Excel exports in many locales)
cols_ar = ["listing_id", "المدينة", "الإيجار_الشهري", "المساحة_م2", "الغرف", "مفروشة", "تاريخ_الإعلان"]
with io.open(os.path.join(T1, "rental_listings_ar_semicolon.csv"), "w", encoding="utf-8-sig", newline="") as f:
    w = csv.writer(f, delimiter=";")
    w.writerow(cols_ar)
    for r in rentals[:30]:
        w.writerow([r["listing_id"], r["city_ar"], r["monthly_rent"],
                    str(r["area_m2"]).replace(".", ","), r["rooms"],
                    "نعم" if r["furnished"] == "yes" else "لا", r["listed_on"]])

# Excel: a title block above the header, serial dates in General format on the Raw sheet
try:
    import openpyxl
    from openpyxl.utils import get_column_letter
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Raw"
    ws["A1"] = "Rental listings - export (do not edit)"
    ws["A2"] = "Exported 2026-05-10"
    hdr = ["listing_id", "city", "monthly_rent", "area_m2", "rooms", "furnished", "listed_on"]
    for j, h in enumerate(hdr, 1):
        ws.cell(row=4, column=j, value=h)
    epoch = dt.date(1899, 12, 30)
    for i, r in enumerate(rentals[:40], 5):
        ws.cell(row=i, column=1, value=r["listing_id"])
        ws.cell(row=i, column=2, value=r["city"])
        ws.cell(row=i, column=3, value=r["monthly_rent"] if r["monthly_rent"] != "" else None)
        ws.cell(row=i, column=4, value=r["area_m2"] if r["area_m2"] != "" else None)
        ws.cell(row=i, column=5, value=r["rooms"] if r["rooms"] != "" else None)
        ws.cell(row=i, column=6, value=r["furnished"])
        d = dt.date.fromisoformat(r["listed_on"])
        c = ws.cell(row=i, column=7, value=(d - epoch).days)     # a serial number, General format
        c.number_format = "General"
    ws2 = wb.create_sheet("Clean")
    for j, h in enumerate(hdr, 1):
        ws2.cell(row=1, column=j, value=h)
    for i, r in enumerate(rentals[:40], 2):
        ws2.cell(row=i, column=1, value=r["listing_id"])
        ws2.cell(row=i, column=2, value=r["city"])
        ws2.cell(row=i, column=3, value=r["monthly_rent"] if r["monthly_rent"] != "" else None)
        ws2.cell(row=i, column=4, value=r["area_m2"] if r["area_m2"] != "" else None)
        ws2.cell(row=i, column=5, value=r["rooms"] if r["rooms"] != "" else None)
        ws2.cell(row=i, column=6, value=r["furnished"])
        c = ws2.cell(row=i, column=7, value=dt.date.fromisoformat(r["listed_on"]))
        c.number_format = "yyyy-mm-dd"
    wb.save(os.path.join(T1, "rental_listings.xlsx"))
    HAVE_XLSX = True
except ImportError:
    HAVE_XLSX = False

# the school-registry API, as three pages of fixtures + the whole document
LEVELS = ["Primary", "Preparatory", "Secondary"]
schools = []
for i in range(1, 14):
    city_en, city_ar, _, _ = CITIES[(i - 1) % len(CITIES)]
    boys, girls = rnd.randint(80, 320), rnd.randint(80, 320)
    s = {
        "id": 200 + i,
        "name": {"en": "School %d (%s)" % (i, city_en), "ar": "مدرسة %d (%s)" % (i, city_ar)},
        "governorate": GOV[city_en],
        "level": LEVELS[i % 3],
        "enrolment": {"boys": boys, "girls": girls},
        "location": {"lat": round(31.3 + rnd.random() * 0.3, 4), "lon": round(34.2 + rnd.random() * 0.2, 4)},
        "facilities": rnd.sample(["library", "lab", "playground", "canteen", "computer room"], rnd.randint(0, 3)),
    }
    if i % 4 != 0:
        s["contacts"] = {"phone": "08-%07d" % rnd.randint(1000000, 9999999)}
        if i % 3 == 0:
            s["contacts"]["email"] = "school%d@example.org" % i
    schools.append(s)
per = 5
pages = [schools[k:k + per] for k in range(0, len(schools), per)]
for n, chunk in enumerate(pages, 1):
    doc = {"page": n, "per_page": per, "total": len(schools), "pages": len(pages),
           "next": ("https://api.example.org/schools?page=%d" % (n + 1)) if n < len(pages) else None,
           "results": chunk}
    with io.open(os.path.join(T1, "schools_api_page%d.json" % n), "w", encoding="utf-8", newline="") as f:
        json.dump(doc, f, ensure_ascii=False, indent=2)
with io.open(os.path.join(T1, "schools.json"), "w", encoding="utf-8", newline="") as f:
    json.dump({"results": schools}, f, ensure_ascii=False, indent=2)

fact("T1 rental_listings.csv: %d rows; %d blank area_m2, %d blank monthly_rent, %d blank rooms; columns %s"
     % (len(rentals), 4, 2, 1, ", ".join(cols)))
fact("T1 rental_listings_ar_semicolon.csv: 30 rows, ';' delimiter, decimal comma in المساحة_م2, UTF-8 with BOM")
fact("T1 rental_listings.xlsx: sheet 'Raw' has 3 title rows above the header at row 4 and listed_on as serial numbers "
     "(General format); sheet 'Clean' has real dates. 40 rows. %s" % ("" if HAVE_XLSX else "(openpyxl missing: not written!)"))
fact("T1 schools: %d schools in %d API pages of %d (schools_api_page1..%d.json) and one document schools.json; "
     "contacts missing on %d schools, email present on %d"
     % (len(schools), len(pages), per, len(pages),
        sum(1 for s in schools if "contacts" not in s),
        sum(1 for s in schools if "email" in s.get("contacts", {}))))

# ------------------------------------------------------------------ Practical 1: cafeteria

rnd = random.Random(1103)
MENU = [
    ("Espresso", "Drink", 6.0), ("Latte", "Drink", 9.0), ("Tea", "Drink", 3.0), ("Karak Tea", "Drink", 5.0),
    ("Orange Juice", "Drink", 8.0), ("Water 500ml", "Drink", 2.0),
    ("Falafel Sandwich", "Food", 7.0), ("Shawarma Wrap", "Food", 12.0), ("Cheese Manakish", "Food", 8.0),
    ("Chicken Rice", "Food", 15.0), ("Lentil Soup", "Food", 6.0),
    ("Chips", "Snack", 4.0), ("Date Cookie", "Snack", 3.5), ("Nuts Cup", "Snack", 5.0),
    ("Baklava", "Dessert", 6.5), ("Kunafa Slice", "Dessert", 9.0),
]
WEIGHT = {"Espresso": 9, "Latte": 7, "Tea": 14, "Karak Tea": 8, "Orange Juice": 6, "Water 500ml": 12,
          "Falafel Sandwich": 13, "Shawarma Wrap": 9, "Cheese Manakish": 8, "Chicken Rice": 5, "Lentil Soup": 3,
          "Chips": 8, "Date Cookie": 5, "Nuts Cup": 3, "Baklava": 3, "Kunafa Slice": 2}
HOUR_W = {8: 6, 9: 9, 10: 12, 11: 14, 12: 22, 13: 20, 14: 11, 15: 6, 16: 3}
PAY = ["Cash"] * 5 + ["Card"] * 2 + ["Wallet"] * 2 + ["Voucher"]
CUST = ["Student"] * 7 + ["Staff"] * 2 + ["Visitor"]
menu_price = {m[0]: m[2] for m in MENU}
menu_cat = {m[0]: m[1] for m in MENU}

items, wts = zip(*WEIGHT.items())
hours, hw = zip(*HOUR_W.items())
orders = []
oid = 5000
day = dt.date(2026, 9, 1)
while day.month == 9:
    if day.weekday() in (4, 5):            # Friday, Saturday: closed
        day += dt.timedelta(days=1)
        continue
    n = rnd.randint(22, 34) + (6 if day.weekday() == 6 else 0)
    for _ in range(n):
        oid += 1
        h = rnd.choices(hours, hw)[0]
        t = dt.datetime(day.year, day.month, day.day, h, rnd.randint(0, 59), rnd.randint(0, 59))
        for _line in range(rnd.choice([1, 1, 1, 2])):
            item = rnd.choices(items, wts)[0]
            qty = rnd.choice([1, 1, 1, 2, 2, 3])
            orders.append({"order_id": oid, "order_time": t, "item": item, "category": menu_cat[item],
                           "qty": qty, "unit_price": menu_price[item], "payment": rnd.choice(PAY),
                           "customer_type": rnd.choice(CUST)})
    day += dt.timedelta(days=1)

# --- deliberate mess ---------------------------------------------------------
clean_rows = len(orders)
# 1. duplicated lines (double-click on the till)
dups = rnd.sample(range(len(orders)), 14)
for k in sorted(dups, reverse=True):
    orders.insert(k + 1, dict(orders[k]))
# 2. missing payment method
miss = rnd.sample(range(len(orders)), 17)
for k in miss:
    orders[k]["payment"] = ""
# 3. spelling variants of item names
variants = {"Latte": ["latte", "LATTE", "Latte "], "Tea": ["tea", "Tea "], "Espresso": ["espresso", "Espresso "]}
nvar = 0
for o in orders:
    if o["item"] in variants and rnd.random() < 0.06:
        o["item"] = rnd.choice(variants[o["item"]])
        nvar += 1
# 4. a price typo (a decimal point lost) and a zero/negative quantity (refunds)
typo_at = next(i for i, o in enumerate(orders) if o["item"] == "Shawarma Wrap")
orders[typo_at]["unit_price"] = 120.0
ref = rnd.sample([i for i, o in enumerate(orders) if o["qty"] == 1], 5)
for k in ref:
    orders[k]["qty"] = -1
# 5. an unparseable price
bad_at = next(i for i, o in enumerate(orders) if o["item"] == "Tea" and i > typo_at)
orders[bad_at]["unit_price"] = "n/a"

cols = ["order_id", "order_time", "item", "category", "qty", "unit_price", "payment", "customer_type"]
with io.open(os.path.join(P1, "cafeteria_orders.csv"), "w", encoding="utf-8", newline="") as f:
    w = csv.writer(f)
    w.writerow(cols)
    for o in orders:
        w.writerow([o["order_id"], o["order_time"].strftime("%Y-%m-%d %H:%M:%S"), o["item"], o["category"],
                    o["qty"], o["unit_price"], o["payment"], o["customer_type"]])
with io.open(os.path.join(P1, "cafeteria_menu.csv"), "w", encoding="utf-8", newline="") as f:
    w = csv.writer(f)
    w.writerow(["item", "category", "list_price"])
    for m in MENU:
        w.writerow(m)

try:
    import openpyxl
    from openpyxl.worksheet.table import Table, TableStyleInfo
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Orders"
    ws.append(cols)
    for o in orders:
        ws.append([o["order_id"], o["order_time"], o["item"], o["category"], o["qty"], o["unit_price"],
                   o["payment"] or None, o["customer_type"]])
    last = len(orders) + 1
    for row in ws.iter_rows(min_row=2, min_col=2, max_col=2):
        row[0].number_format = "yyyy-mm-dd hh:mm"
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 18
    tab = Table(displayName="tblOrders", ref="A1:H%d" % last)
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(tab)
    ws.freeze_panes = "A2"
    wm = wb.create_sheet("Menu")
    wm.append(["item", "category", "list_price"])
    for m in MENU:
        wm.append(list(m))
    tab2 = Table(displayName="tblMenu", ref="A1:C%d" % (len(MENU) + 1))
    tab2.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    wm.add_table(tab2)
    wb.save(os.path.join(P1, "cafeteria_orders.xlsx"))
except ImportError:
    pass

# --- the facts, computed from what was written -------------------------------
import pandas as pd

df = pd.read_csv(os.path.join(P1, "cafeteria_orders.csv"))
raw_rows = len(df)
df["unit_price_num"] = pd.to_numeric(df["unit_price"], errors="coerce")
df["order_time"] = pd.to_datetime(df["order_time"])
df["line_total"] = df["qty"] * df["unit_price_num"]
dup_lines = int(df.duplicated().sum())
tidy = df.drop_duplicates()
good = tidy[(tidy["qty"] > 0) & (tidy["unit_price_num"] < 50) & tidy["unit_price_num"].notna()].copy()
good["item_clean"] = good["item"].str.strip().str.title()
rev_by_item = good.groupby("item_clean")["line_total"].sum().sort_values(ascending=False)
hour_orders = good.groupby(good["order_time"].dt.hour)["order_id"].nunique().sort_values(ascending=False)
cat_rev = good.groupby("category")["line_total"].sum().sort_values(ascending=False)

fact("P1 cafeteria_orders.csv: %d rows (lines), columns %s; period 2026-09-01..2026-09-30, closed Fri/Sat"
     % (raw_rows, ", ".join(cols)))
fact("P1 mess: %d exactly-duplicated lines; %d blank payment; %d item spelling variants (case/space of Latte, Tea, Espresso); "
     "1 price typo (Shawarma Wrap at 120.0); 1 unit_price 'n/a' (pandas read_csv reads 'n/a' as missing BY DEFAULT, so there the column loads as float64 with one NaN; Excel and Power BI see it as text/error); 5 refunds with qty = -1" % (dup_lines, int(df['payment'].isna().sum()), nvar))
fact("P1 distinct order_id: %d; distinct items in raw: %d; distinct items after strip+title: %d"
     % (df["order_id"].nunique(), df["item"].nunique(), good["item_clean"].nunique()))
fact("P1 after dropping duplicates, refunds, the typo and 'n/a': %d lines; total revenue = %.2f"
     % (len(good), good["line_total"].sum()))
fact("P1 top 3 items by revenue: " + "; ".join("%s=%.2f" % (k, v) for k, v in rev_by_item.head(3).items()))
fact("P1 revenue by category: " + "; ".join("%s=%.2f" % (k, v) for k, v in cat_rev.items()))
fact("P1 busiest hour by distinct orders: %d:00 (%d orders); then %s"
     % (hour_orders.index[0], hour_orders.iloc[0], ", ".join("%d:00 (%d)" % (h, n) for h, n in hour_orders.iloc[1:3].items())))

with io.open(os.path.join(ROOT, "docs", "datasets.md"), "w", encoding="utf-8", newline="\n") as f:
    f.write("# Shared sample datasets\n\n"
            "Generated by `python tools/make_datasets.py` (deterministic). Lessons link to the files under\n"
            "`downloads/` with `href=\"@/downloads/<chapter>/<file>\"` and quote the numbers below — never\n"
            "invent a different number for the same data. Re-run the script, never hand-edit the files.\n\n"
            "## Facts (computed from the files)\n\n")
    for s in facts:
        f.write("- " + s + "\n")
    f.write("\n## Files\n\n"
            "| File | What it is |\n|---|---|\n"
            "| `downloads/t1/rental_listings.csv` | 60 listings, comma CSV, UTF-8, ISO dates, a few blanks |\n"
            "| `downloads/t1/rental_listings_ar_semicolon.csv` | 30 rows, `;` delimiter, decimal comma, Arabic headers and city names, UTF-8 BOM |\n"
            "| `downloads/t1/rental_listings.xlsx` | sheet `Raw` (title rows, header on row 4, serial-number dates) and `Clean` |\n"
            "| `downloads/t1/schools_api_page1..3.json` | a paginated API reply (`results`, `page`, `next`, `total`) |\n"
            "| `downloads/t1/schools.json` | the whole registry as one nested document |\n"
            "| `downloads/p1/cafeteria_orders.csv` | one month of cafeteria order lines (messy on purpose) |\n"
            "| `downloads/p1/cafeteria_orders.xlsx` | the same as an Excel workbook: Table `tblOrders`, sheet `Menu` with Table `tblMenu` |\n"
            "| `downloads/p1/cafeteria_menu.csv` | the menu: item, category, list_price |\n")
import sys
sys.stdout.reconfigure(encoding="utf-8")
print("\n".join(facts))
